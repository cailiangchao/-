from flask import Blueprint, render_template_string, request, jsonify, send_file
import csv
import io
import json
import os
from datetime import datetime

# 创建Blueprint
search_bp = Blueprint('search', __name__)

# 定义HTML模板头部
HTML_HEAD = '''<!doctype html>
<html lang="zh">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>设备数据高级查询</title>
    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css">
    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/bootstrap-icons@1.11.0/font/bootstrap-icons.css">
    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/flatpickr/dist/flatpickr.min.css">
'''

# 定义CSS样式
CSS_STYLES = '''<style>
    :root {
        --primary-color: #4285f4;
        --secondary-color: #34a853;
        --danger-color: #ea4335;
        --light-gray: #f5f5f5;
        --dark-gray: #333;
        --glass-bg: rgba(255, 255, 255, 0.15);
    }
    
    body {
        font-family: 'Microsoft YaHei', '微软雅黑', Arial, sans-serif;
        background: #ffffff;
        color: var(--dark-gray);
        min-height: 100vh;
        margin: 0;
        padding: 20px;
    }
    
    .container {
        max-width: 1400px;
        padding: 20px;
    }
    
    .card {
        border-radius: 12px;
        box-shadow: 0 8px 32px rgba(0, 0, 0, 0.1);
        margin-bottom: 25px;
        border: 1px solid #e0e0e0;  /* 改为灰色边框 */
        transition: all 0.3s ease;
        overflow: hidden;
        background: white;  /* 改为纯白色背景 */
    }
    

    
    .card-header {
        background-color: white;
        border-bottom: 1px solid rgba(0, 0, 0, 0.05);
        padding: 15px 20px;
        font-weight: 600;
        color: var(--primary-color);
        border-radius: 10px 10px 0 0 !important;
    }
    
    .card-body {
        padding: 20px;
    }
    
    .btn-primary {
        background-color: var(--primary-color);
        border-color: var(--primary-color);
        transition: all 0.2s ease;
        position: relative;
        overflow: hidden;
    }
    
    .btn-primary:hover {
        background-color: #3367d6;
        border-color: #3367d6;
        box-shadow: 0 4px 8px rgba(0, 0, 0, 0.1);
    }
    
    
    
    .btn-outline-primary {
        color: var(--primary-color);
        border-color: var(--primary-color);
        transition: all 0.2s ease;
        position: relative;
        overflow: hidden;
    }
    
    .btn-outline-primary:hover {
        background-color: var(--primary-color);
        color: white;
        box-shadow: 0 4px 8px rgba(0, 0, 0, 0.1);
    }
    
    
    
    .form-control:focus, .form-select:focus {
        border-color: var(--primary-color);
        box-shadow: 0 0 0 0.25rem rgba(66, 133, 244, 0.25);
    }
    
    .table {
        vertical-align: middle;
        border-collapse: separate;
        border-spacing: 0;
    }
    
    .table th {
        background-color: #4285f4;  /* 改为蓝色背景 */
        font-weight: 600;
        color: white;
        padding: 12px 15px;
        border: none;
        position: sticky;
        top: 0;
        backdrop-filter: blur(10px);
    }
    
    .table td {
        padding: 12px 15px;
        border-bottom: 1px solid #ddd;  /* 改为灰色边框 */
        background-color: white;  /* 改为白色背景 */
        transition: all 0.2s ease;
    }
    
    .table tr:hover td {
        background-color: #f5f5f5;  /* 鼠标悬停改为浅灰色 */
        
    }
    
    .pagination .page-item.active .page-link {
        background-color: var(--primary-color);
        border-color: var(--primary-color);
    }
    
    .pagination .page-link {
        color: var(--primary-color);
    }
    
    .search-filters {
        display: flex;
        flex-wrap: wrap;
        gap: 10px;
        margin-bottom: 15px;
    }
    
    .search-filter-tag {
        background-color: rgba(66, 133, 244, 0.1);
        color: var(--primary-color);
        border-radius: 20px;
        padding: 5px 12px;
        font-size: 14px;
        display: inline-flex;
        align-items: center;
        gap: 5px;
    }
    
    .search-filter-tag i {
        cursor: pointer;
    }
    
    .sort-icon {
        cursor: pointer;
        margin-left: 5px;
    }
    
    .sort-icon.active {
        color: var(--primary-color);
    }
    
    .loading-overlay {
        position: fixed;
        top: 0;
        left: 0;
        width: 100%;
        height: 100%;
        background-color: rgba(255, 255, 255, 0.8);
        display: flex;
        justify-content: center;
        align-items: center;
        z-index: 1000;
        visibility: hidden;
        opacity: 0;
        transition: visibility 0s, opacity 0.3s;
    }
    
    .loading-overlay.show {
        visibility: visible;
        opacity: 1;
    }
    
    .spinner-border {
        width: 3rem;
        height: 3rem;
        color: var(--primary-color);
    }
    
    .no-results {
        text-align: center;
        padding: 40px 0;
        color: #666;
    }
    
    .no-results i {
        font-size: 48px;
        color: #ddd;
        margin-bottom: 15px;
        display: block;
    }
    
    .flatpickr-input {
        background-color: white !important;
    }
    
    @media (max-width: 768px) {
        .card-header {
            padding: 12px 15px;
        }
        
        .card-body {
            padding: 15px;
        }
        
        .container {
            padding: 10px;
        }
        
        .table-responsive {
            border: none;
        }
    }
</style>
</head>
'''

# 定义HTML主体部分
HTML_BODY = '''<body>
    <div class="container mt-4">
        <div class="d-flex justify-content-between align-items-center mb-4">
            <h1 class="h3">设备数据高级查询</h1>
            <div>
                <a href="/" class="btn btn-outline-primary me-2">
                    <i class="bi bi-plus-circle"></i> 新增记录
                </a>
                <button id="overdueBtn" class="btn btn-outline-warning me-2">
                    <i class="bi bi-exclamation-triangle"></i> 超期未消毒设备
                </button>
                <button id="exportBtn" class="btn btn-outline-success">
                    <i class="bi bi-file-earmark-excel"></i> 导出结果
                </button>
            </div>
        </div>

        <!-- 统计卡片 -->
        <div class="row mb-4">
            <div class="col-md-3">
                <div class="card">
                    <div class="card-body text-center">
                        <h5 class="card-title">总记录数</h5>
                        <h2 class="card-text" id="totalCount">0</h2>
                    </div>
                </div>
            </div>
            <div class="col-md-3">
                <div class="card">
                    <div class="card-body text-center">
                        <h5 class="card-title">正常运行</h5>
                        <h2 class="card-text text-success" id="normalCount">0</h2>
                    </div>
                </div>
            </div>
            <div class="col-md-3">
                <div class="card">
                    <div class="card-body text-center">
                        <h5 class="card-title">故障设备</h5>
                        <h2 class="card-text text-danger" id="errorCount">0</h2>
                    </div>
                </div>
            </div>
            <div class="col-md-3">
                <div class="card">
                    <div class="card-body text-center">
                        <h5 class="card-title">空闲设备</h5>
                        <h2 class="card-text text-warning" id="idleCount">0</h2>
                    </div>
                </div>
            </div>
        </div>
        
        <div class="card">
            <div class="card-header d-flex justify-content-between align-items-center">
                <span>查询条件</span>
                <button id="toggleFilters" class="btn btn-sm btn-outline-primary">
                    <i class="bi bi-funnel"></i> 显示/隐藏筛选器
                </button>
            </div>
            <div class="card-body" id="filterSection">
                <form id="searchForm" class="row g-3">
                    <div class="col-md-4">
                        <label for="name" class="form-label">设备名称</label>
                        <select id="name" name="name" class="form-select">
                            <option value="">全部</option>
                            <option value="暖箱">暖箱</option>
                            <option value="长颈鹿多功能暖箱">长颈鹿多功能暖箱</option>
                            <option value="辐射台">辐射台</option>
                            <option value="菲萍呼吸机">菲萍呼吸机</option>
                            <option value="SLE呼吸机">SLE呼吸机</option>
                            <option value="千禧呼吸机">千禧呼吸机</option>
                            <option value="E360呼吸机">E360呼吸机</option>
                        </select>
                    </div>
                    
                    <div class="col-md-4">
                        <label for="number" class="form-label">设备编号</label>
                        <input type="text" id="number" name="number" class="form-control" placeholder="输入设备编号">
                    </div>
                    
                    <div class="col-md-4">
                        <label for="status" class="form-label">运行状态</label>
                        <select id="status" name="status" class="form-select">
                            <option value="">全部</option>
                            <option value="正常运行">正常运行</option>
                            <option value="故障">故障</option>
                            <option value="空闲">空闲</option>
                        </select>
                    </div>
                    
                    <div class="col-md-4">
                        <label for="user" class="form-label">使用人</label>
                        <input type="text" id="user" name="user" class="form-control" placeholder="输入使用人姓名">
                    </div>
                    
                    <div class="col-md-4">
                        <label for="bed" class="form-label">床号</label>
                        <input type="text" id="bed" name="bed" class="form-control" placeholder="输入床号">
                    </div>

                    <div class="col-md-4">
                        <label for="location" class="form-label">放置位置</label>
                        <input type="text" id="location" name="location" class="form-control" placeholder="输入放置位置">
                    </div>
                    
                    <div class="col-md-4">
                        <label for="person_name" class="form-label">患儿姓名</label>
                        <input type="text" id="person_name" name="person_name" class="form-control" placeholder="输入患儿姓名">
                    </div>
                    
                    <div class="col-md-4">
                        <label for="patient_id" class="form-label">住院号</label>
                        <input type="text" id="patient_id" name="patient_id" class="form-control" placeholder="输入住院号">
                    </div>
                    
                    <div class="col-md-4">
                        <label for="disinfector" class="form-label">消毒人</label>
                        <input type="text" id="disinfector" name="disinfector" class="form-control" placeholder="输入消毒人姓名">
                    </div>
                    
                    <div class="col-md-4">
                        <label for="notes" class="form-label">备注</label>
                        <input type="text" id="notes" name="notes" class="form-control" placeholder="输入备注关键词">
                    </div>

                    <div class="col-md-4">
                        <label for="overdue_disinfect" class="form-label">终末消毒超期</label>
                        <select id="overdue_disinfect" name="overdue_disinfect" class="form-select">
                            <option value="">全部</option>
                            <option value="7">≥7天未消毒</option>
                        </select>
                    </div>
                    
                    <div class="col-12 mt-4">
                        <div class="d-flex justify-content-between">
                            <button type="submit" class="btn btn-primary">
                                <i class="bi bi-search"></i> 查询
                            </button>
                            <button type="reset" class="btn btn-outline-secondary">
                                <i class="bi bi-arrow-counterclockwise"></i> 重置
                            </button>
                        </div>
                    </div>
                </form>
            </div>
        </div>
        
        <div class="card">
            <div class="card-header d-flex justify-content-between align-items-center">
                <span>查询结果</span>
                <div class="d-flex align-items-center">
                    <span class="me-3" id="resultCount">共 <strong>0</strong> 条记录</span>
                    <select id="pageSize" class="form-select form-select-sm" style="width: auto;">
                        <option value="10">10条/页</option>
                        <option value="20">20条/页</option>
                        <option value="50">50条/页</option>
                        <option value="100">100条/页</option>
                    </select>
                </div>
            </div>
            <div class="card-body p-0">
                <div id="activeFilters" class="px-3 pt-3 search-filters"></div>
                
                <div class="table-responsive">
                    <table class="table table-hover mb-0" id="resultsTable">
                        <thead>
                            <tr>
                                <th data-field="设备名称">设备名称 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="设备编号">设备编号 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="使用日期">使用日期 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="运行状态">运行状态 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="使用人">使用人 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="床号">床号 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="放置位置">放置位置 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="患儿姓名">患儿姓名 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="住院号">住院号 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="结束日期">结束日期 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="终末消毒日期">终末消毒日期 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="消毒人">消毒人 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                                <th data-field="备注">备注</th>
                                <th data-field="记录时间">记录时间 <i class="bi bi-arrow-down-up sort-icon"></i></th>
                            </tr>
                        </thead>
                        <tbody id="resultsBody">
                            <!-- 结果将通过JavaScript动态填充 -->
                        </tbody>
                    </table>
                </div>
                
                <div id="noResults" class="no-results" style="display: none;">
                    <i class="bi bi-search"></i>
                    <p>没有找到匹配的记录</p>
                    <p class="text-muted">请尝试调整查询条件</p>
                </div>
                
                <div class="d-flex justify-content-between align-items-center p-3">
                    <div id="pageInfo">显示 0-0 条，共 0 条</div>
                    <nav aria-label="Page navigation">
                        <ul class="pagination mb-0" id="pagination">
                            <!-- 分页将通过JavaScript动态填充 -->
                        </ul>
                    </nav>
                </div>
            </div>
        </div>
    </div>
    
    <div class="loading-overlay" id="loadingOverlay">
        <div class="spinner-border" role="status">
            <span class="visually-hidden">加载中...</span>
        </div>
    </div>
</body>
</html>
'''

# 定义JavaScript部分
JS_SCRIPT = '''<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script>
    <script src="https://cdn.jsdelivr.net/npm/flatpickr"></script>
    <script src="https://cdn.jsdelivr.net/npm/flatpickr/dist/l10n/zh.js"></script>
    <script>
        document.addEventListener('DOMContentLoaded', function() {
            // 初始化日期选择器
            flatpickr(".date-picker", {
                locale: "zh",
                dateFormat: "Y-m-d",
                allowInput: true
            });
            
            // 显示/隐藏筛选器
            document.getElementById('toggleFilters').addEventListener('click', function() {
                const filterSection = document.getElementById('filterSection');
                if (filterSection.style.display === 'none') {
                    filterSection.style.display = 'block';
                } else {
                    filterSection.style.display = 'none';
                }
            });
            
            // 全局变量
            let allData = [];
            let currentPage = 1;
            let pageSize = 10;
            let sortField = '';
            let sortDirection = 'asc';
            let activeFilters = {};
            
            // 获取数据
            function fetchData() {
                showLoading();
                
                // 收集表单数据
                const formData = new FormData(document.getElementById('searchForm'));
                const searchParams = new URLSearchParams();
                
                for (const [key, value] of formData.entries()) {
                    if (value) {
                        searchParams.append(key, value);
                    }
                }
                
                // 发送请求
                fetch('/api/search?' + searchParams.toString())
                    .then(response => response.json())
                    .then(data => {
                        allData = data;
                        updateActiveFilters(formData);
                        updateStats(data);
                        renderTable();
                        hideLoading();
                    })
                    .catch(error => {
                        console.error('Error fetching data:', error);
                        hideLoading();
                        alert('获取数据失败，请重试');
                    });
            }

            // 更新统计数据
            function updateStats(data) {
                document.getElementById('totalCount').textContent = data.length;
                document.getElementById('normalCount').textContent = 
                    data.filter(item => item['运行状态'] === '正常运行').length;
                document.getElementById('errorCount').textContent = 
                    data.filter(item => item['运行状态'] === '故障').length;
                document.getElementById('idleCount').textContent = 
                    data.filter(item => item['运行状态'] === '空闲').length;
            }
            
            // 更新活动筛选器标签
            function updateActiveFilters(formData) {
                activeFilters = {};
                const filterLabels = {
                    'name': '设备名称',
                    'number': '设备编号',
                    'status': '运行状态',
                    'user': '使用人',
                    'bed': '床号',
                    'location': '放置位置',
                    'person_name': '患儿姓名',
                    'patient_id': '住院号',
                    'disinfector': '消毒人',
                    'notes': '备注',
                    'overdue_disinfect': '终末消毒超期' 
                };
                
                for (const [key, value] of formData.entries()) {
                    if (value && filterLabels[key]) {
                        activeFilters[key] = {
                            label: filterLabels[key],
                            value: value
                        };
                    }
                }
                
                renderActiveFilters();
            }
            
            // 渲染活动筛选器
            function renderActiveFilters() {
                const container = document.getElementById('activeFilters');
                container.innerHTML = '';
                
                if (Object.keys(activeFilters).length === 0) {
                    container.style.display = 'none';
                    return;
                }
                
                container.style.display = 'flex';
                
                for (const [key, filter] of Object.entries(activeFilters)) {
                    const tag = document.createElement('div');
                    tag.className = 'search-filter-tag';
                    tag.innerHTML = `
                        <span>${filter.label}: ${filter.value}</span>
                        <i class="bi bi-x-circle" data-filter-key="${key}"></i>
                    `;
                    container.appendChild(tag);
                    
                    // 添加删除筛选器的事件
                    tag.querySelector('i').addEventListener('click', function() {
                        const filterKey = this.getAttribute('data-filter-key');
                        document.getElementById(filterKey).value = '';
                        fetchData();
                    });
                }
            }
            
            // 渲染表格
            function renderTable() {
                const tbody = document.getElementById('resultsBody');
                tbody.innerHTML = '';
                
                // 应用排序
                let displayData = [...allData];
                if (sortField) {
                    displayData.sort((a, b) => {
                        const valA = a[sortField] || '';
                        const valB = b[sortField] || '';
                        return sortDirection === 'asc' 
                            ? valA.localeCompare(valB) 
                            : valB.localeCompare(valA);
                    });
                }
                
                // 应用分页
                const startIndex = (currentPage - 1) * pageSize;
                const endIndex = Math.min(startIndex + pageSize, displayData.length);
                const paginatedData = displayData.slice(startIndex, endIndex);
                
                // 更新结果计数
                document.getElementById('resultCount').innerHTML = 
                    `共 <strong>${displayData.length}</strong> 条记录`;
                document.getElementById('pageInfo').textContent = 
                    `显示 ${startIndex + 1}-${endIndex} 条，共 ${displayData.length} 条`;
                
                // 渲染表格行
                if (paginatedData.length === 0) {
                    document.getElementById('noResults').style.display = 'block';
                    document.getElementById('resultsTable').style.display = 'none';
                } else {
                    document.getElementById('noResults').style.display = 'none';
                    document.getElementById('resultsTable').style.display = 'table';
                    
                    paginatedData.forEach(item => {
                        const row = document.createElement('tr');
                        row.innerHTML = `
                            <td>${item['设备名称'] || ''}</td>
                            <td>${item['设备编号'] || ''}</td>
                            <td>${item['使用日期'] || ''}</td>
                            <td>
                                ${item['运行状态'] || ''}
                                ${item['运行状态'] === '正常运行' ? '✅' : ''}
                                ${item['运行状态'] === '故障' ? '❌' : ''}
                                ${item['运行状态'] === '空闲' ? '⏸️' : ''}
                            </td>
                            <td>${item['使用人'] || ''}</td>
                            <td>${item['床号'] || ''}</td>
                            <td>${item['放置位置'] || ''}</td>
                            <td>${item['患儿姓名'] || ''}</td>
                            <td>${item['住院号'] || ''}</td>
                            <td>${item['结束日期'] || ''}</td>
                            <td>${item['终末消毒日期'] || ''}</td>
                            <td>${item['消毒人'] || ''}</td>
                            <td>${item['备注'] || ''}</td>
                            <td>${item['记录时间'] || ''}</td>
                        `;
                        tbody.appendChild(row);
                    });
                }
                
                // 渲染分页
                renderPagination(displayData.length);
            }
            
            // 渲染分页控件
            function renderPagination(totalItems) {
                const totalPages = Math.ceil(totalItems / pageSize);
                const pagination = document.getElementById('pagination');
                pagination.innerHTML = '';
                
                // 上一页按钮
                const prevLi = document.createElement('li');
                prevLi.className = `page-item ${currentPage === 1 ? 'disabled' : ''}`;
                prevLi.innerHTML = `<a class="page-link" href="#">上一页</a>`;
                prevLi.addEventListener('click', (e) => {
                    e.preventDefault();
                    if (currentPage > 1) {
                        currentPage--;
                        renderTable();
                    }
                });
                pagination.appendChild(prevLi);
                
                // 页码按钮
                for (let i = 1; i <= totalPages; i++) {
                    const pageLi = document.createElement('li');
                    pageLi.className = `page-item ${i === currentPage ? 'active' : ''}`;
                    pageLi.innerHTML = `<a class="page-link" href="#">${i}</a>`;
                    pageLi.addEventListener('click', (e) => {
                        e.preventDefault();
                        currentPage = i;
                        renderTable();
                    });
                    pagination.appendChild(pageLi);
                }
                
                // 下一页按钮
                const nextLi = document.createElement('li');
                nextLi.className = `page-item ${currentPage === totalPages ? 'disabled' : ''}`;
                nextLi.innerHTML = `<a class="page-link" href="#">下一页</a>`;
                nextLi.addEventListener('click', (e) => {
                    e.preventDefault();
                    if (currentPage < totalPages) {
                        currentPage++;
                        renderTable();
                    }
                });
                pagination.appendChild(nextLi);
            }
            
            // 排序处理
            document.querySelectorAll('.sort-icon').forEach(icon => {
                icon.addEventListener('click', function() {
                    const field = this.parentElement.getAttribute('data-field');
                    
                    // 重置所有排序图标
                    document.querySelectorAll('.sort-icon').forEach(i => {
                        i.classList.remove('active');
                        i.classList.remove('bi-arrow-up');
                        i.classList.remove('bi-arrow-down');
                    });
                    
                    // 设置当前排序状态
                    if (sortField === field) {
                        sortDirection = sortDirection === 'asc' ? 'desc' : 'asc';
                    } else {
                        sortField = field;
                        sortDirection = 'asc';
                    }
                    
                    // 更新排序图标
                    this.classList.add('active');
                    this.classList.add(sortDirection === 'asc' ? 'bi-arrow-up' : 'bi-arrow-down');
                    
                    renderTable();
                });
            });
            
            // 分页大小变更
            document.getElementById('pageSize').addEventListener('change', function() {
                pageSize = parseInt(this.value);
                currentPage = 1;
                renderTable();
            });
            
            // 超期未消毒设备按钮
            document.getElementById('overdueBtn').addEventListener('click', function() {
                // 重置表单
                document.getElementById('searchForm').reset();
                
                // 设置超期筛选条件
                document.getElementById('overdue_disinfect').value = '7';
                
                // 重置页码并查询
                currentPage = 1;
                fetchData();
            });

            // 导出按钮
            document.getElementById('exportBtn').addEventListener('click', function() {
                if (allData.length === 0) {
                    alert('没有数据可导出');
                    return;
                }
                
                showLoading();
                
                // 发送导出请求
                fetch('/api/export', {
                    method: 'POST',
                    headers: {
                        'Content-Type': 'application/json',
                    },
                    body: JSON.stringify({
                        data: allData,
                        filters: activeFilters
                    })
                })
                .then(response => response.blob())
                .then(blob => {
                    const url = window.URL.createObjectURL(blob);
                    const a = document.createElement('a');
                    a.href = url;
                    a.download = `设备数据_${new Date().toISOString().slice(0, 10)}.xlsx`;
                    document.body.appendChild(a);
                    a.click();
                    window.URL.revokeObjectURL(url);
                    document.body.removeChild(a);
                    hideLoading();
                })
                .catch(error => {
                    console.error('导出失败:', error);
                    hideLoading();
                    alert('导出失败，请重试');
                });
            });
            
            // 显示加载中
            function showLoading() {
                document.getElementById('loadingOverlay').classList.add('show');
            }
            
            // 隐藏加载中
            function hideLoading() {
                document.getElementById('loadingOverlay').classList.remove('show');
            }
            
            // 表单提交
            document.getElementById('searchForm').addEventListener('submit', function(e) {
                e.preventDefault();
                currentPage = 1;
                fetchData();
            });
            
            // 初始加载数据
            fetchData();
        });
    </script>
</body>
</html>
'''



@search_bp.route('/search')
def search_page():
    """显示搜索页面"""
    return render_template_string(SEARCH_HTML)

@search_bp.route('/api/search')
def api_search():
    """API: 查询数据"""
    CSV_FILE = 'data.csv'
    
    # 读取所有数据
    if not os.path.isfile(CSV_FILE):
        return jsonify([])
    
    with open(CSV_FILE, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        data = list(reader)
    
    # 获取查询参数并过滤
    filters = {
        'name': request.args.get('name', ''),
        'number': request.args.get('number', ''),
        'status': request.args.get('status', ''),
        'user': request.args.get('user', ''),
        'bed': request.args.get('bed', ''),
        'location': request.args.get('location', ''),
        'person_name': request.args.get('person_name', ''),
        'patient_id': request.args.get('patient_id', ''),
        'disinfector': request.args.get('disinfector', ''),
        'notes': request.args.get('notes', '')
    }
    
    # 应用过滤条件
    filtered_data = data
    for key, value in filters.items():
        if value:
            field_map = {
                'name': '设备名称',
                'number': '设备编号',
                'status': '运行状态',
                'user': '使用人',
                'bed': '床号',
                'location': '放置位置',
                'person_name': '患儿姓名',
                'patient_id': '住院号',
                'disinfector': '消毒人',
                'notes': '备注'
            }
            field = field_map.get(key)
            if field:
                filtered_data = [row for row in filtered_data 
                               if value == row.get(field, '')]
    #添加终末消毒超期筛选            
    overdue_days = request.args.get('overdue_disinfect', '')
    if overdue_days:
        from datetime import datetime, timedelta
        today = datetime.now()
        threshold_days = int(overdue_days)
        
        def is_overdue(row):
            disinfect_date_str = row.get('终末消毒日期', '')
            if not disinfect_date_str:
                return False  # 没有消毒日期的不算
            
            try:
                disinfect_date = datetime.strptime(disinfect_date_str, '%Y-%m-%d')
                days_diff = (today - disinfect_date).days
                return days_diff >= threshold_days
            except:
                return False
        
        filtered_data = [row for row in filtered_data if is_overdue(row)]
    
    return jsonify(filtered_data)

@search_bp.route('/api/export', methods=['POST'])
def api_export():
    """API: 导出数据为 Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        return jsonify({'error': '缺少 openpyxl 库，请安装: pip install openpyxl'}), 500
    
    data = request.json.get('data', [])
    
    if not data:
        return jsonify({'error': '没有数据可导出'}), 400
    
    # 创建工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "设备数据"
    
    # 写入表头
    headers = list(data[0].keys()) if data else []
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="4285f4", end_color="4285f4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    # 写入数据
    for row_idx, item in enumerate(data, 2):
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col_idx, value=item.get(header, ''))
    
    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'设备数据_{datetime.now().strftime("%Y%m%d")}.xlsx'
    )

# 组合完整HTML模板
SEARCH_HTML = HTML_HEAD + CSS_STYLES + HTML_BODY + JS_SCRIPT